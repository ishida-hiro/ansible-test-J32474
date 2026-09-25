# -*- coding: utf-8 -*-
"""Generate docs/06_検証記録様式.xlsx (blank record sheet for build verification).

作業区分・結果の凡例などの定義は docs/06_検証記録様式.md を正とし、
「記入要領」シートはその表から生成する。
"""
import io, os, sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from md2xlsx import md_tables, finish, TITLE_F, HDR_F, HDR_FILL, BORDER, WRAP  # noqa: E402

SRC = 'docs/06_検証記録様式.md'
DST = 'docs/06_検証記録様式.xlsx'

INPUT_FILL = PatternFill('solid', fgColor='FFF9E5')   # 記入欄
CALC_FILL  = PatternFill('solid', fgColor='F2F2F2')   # 自動計算欄
NOTE_F     = Font(color='595959', size=9)

WORK_KINDS = ['設計書の読み解き', '変数化', 'コード修正', '実行', '確認', '手戻り']
RESULTS    = ['◎', '○', '△', '×', '−']
METHODS    = ['変数のみ', 'ロール修正', '新規ロール', '手動', 'クラウド側', '対象外']
DIFF_KINDS = ['変数のみ', 'ロール修正', '新規ロール', '設計書要確認']
BLANK_ROWS = 60

# 設計書テンプレート（「1.概要」と詳細シート）の項目。節番号は設計書ごとに異なるため記入欄とする。
DESIGN_ITEMS = [
    ('概要', 'OSバージョン', 'OS 版数・エディション'),
    ('概要', 'ドメイン参加', 'ドメイン / ワークグループ'),
    ('概要', 'ユーザアカウント', 'ドメインユーザ・ローカルユーザ'),
    ('概要', 'ドライブ構成', 'C: のサイズ、データ領域のドライブ'),
    ('概要', 'ネットワーク設定', 'IPv6 無効化、IP / DNS'),
    ('概要', '言語設定', '日本語、タイムゾーン'),
    ('概要', 'リモート接続設定', 'リモートデスクトップ、NLA、許可するユーザ'),
    ('概要', '時刻同期設定', 'NTP / ドメイン階層での同期'),
    ('概要', '役割と機能', '追加する役割・機能'),
    ('概要', 'Windows Defender ファイアウォール', '無効化（NSG / セキュリティグループで制御）'),
    ('概要', 'Windows Defender', 'ポリシーで無効化、無効化の時期'),
    ('概要', 'ユーザアカウント制御', 'UAC 無効化'),
    ('概要', 'タスクスケジューラ', 'タスクの登録'),
    ('概要', 'Windows 更新プログラムの適用', '構築時の適用、自動更新の構成 = 無効'),
    ('概要', 'サービス設定', '不要なサービスの無効化'),
    ('概要', 'FTP設定', 'FTP サイト'),
    ('概要', 'ホスト名の戻し', '移行後のホスト名変更'),
    ('詳細', 'サーバーマネージャー', 'コンピュータ名、IE ESC、診断データ、リモート管理'),
    ('詳細', 'Windowsログ', 'イベントログの最大サイズ'),
    ('詳細', 'システムのプロパティ', '仮想メモリ、起動と回復'),
    ('詳細', 'ローカルユーザ / ローカルグループ', 'ユーザ作成、グループ所属、Guest'),
    ('詳細', 'ディスクの管理', 'パーティションスタイル、ファイルシステム、ラベル'),
    ('詳細', '導入ソフトウェア', 'ミドルウェア・ランタイムの導入'),
    ('詳細', 'ネットワーク設定（NIC）', 'バインド、DNS サフィックス、NetBIOS'),
    ('詳細', 'hosts', 'hosts エントリ'),
    ('詳細', '役割と機能の追加（一覧）', '既定で導入済みの項目を含む'),
    ('詳細', 'UAC の無効化（レジストリ）', 'EnableLUA / luafv'),
    ('詳細', 'WindowsUpdate 適用パッチ一覧', '適用済みパッチの記録'),
    ('詳細', 'ローカルグループポリシー', 'Defender / 自動更新'),
]


def dv_list(ws, values, ref):
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    dv.error = '一覧から選んでください'
    dv.errorTitle = '入力値'
    ws.add_data_validation(dv)
    dv.add(ref)


def title(ws, text, note=None):
    ws['A1'] = text
    ws['A1'].font = TITLE_F
    if note:
        ws['A2'] = note
        ws['A2'].font = NOTE_F


def table(ws, header_row, headers, widths, rows=BLANK_ROWS, fixed=None, input_cols=None):
    """ヘッダと記入用の空行を作る。fixed は先頭から埋める固定値の行リスト。"""
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    fixed = fixed or []
    n = max(rows, len(fixed))
    for r in range(n):
        row = header_row + 1 + r
        vals = fixed[r] if r < len(fixed) else []
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=c, value=vals[c - 1] if c - 1 < len(vals) else None)
            cell.border = BORDER
            cell.alignment = WRAP
            if input_cols and c in input_cols:
                cell.fill = INPUT_FILL
    finish(ws, widths, header_row)
    return header_row + 1, header_row + n


# ---------------------------------------------------------------- 各シート
def sheet_basic(wb, ref):
    ws = wb.active
    ws.title = '基本情報'
    title(ws, '検証記録', '黄色の欄に記入します。灰色の欄は他のシートから自動で計算されます。')
    items = [
        ('記入者', ''), ('記入日', ''), ('案件・設計書', ''), ('対象ホスト名', ''), ('用途', ''),
        ('OS', ''), ('環境', ''), ('構築方式', ''), ('何台目', ''), ('作業期間', ''),
        ('Ansible の経験', ''), ('Playbook 実行回数', ''), ('突合せ項目数（99_verify）', ''),
        ('突合せ一致数（99_verify）', ''), ('備考（対象外作業の時間など）', ''),
    ]
    r = 4
    for label, _ in items:
        ws.cell(row=r, column=1, value=label).font = HDR_F
        ws.cell(row=r, column=1).fill = HDR_FILL
        c = ws.cell(row=r, column=2)
        c.fill = INPUT_FILL
        for col in (1, 2):
            ws.cell(row=r, column=col).border = BORDER
            ws.cell(row=r, column=col).alignment = WRAP
        ref[label] = f'B{r}'
        r += 1
    dv_list(ws, ['Azure', 'AWS', 'オンプレミス', 'その他'], ref['環境'])
    dv_list(ws, ['Ansible（既存資産の流用）', 'Ansible（AIで一から生成）', 'Ansible（その他）', '手動'], ref['構築方式'])
    dv_list(ws, ['1台目', '2台目', '3台目以降'], ref['何台目'])
    dv_list(ws, ['なし', '少し', 'あり'], ref['Ansible の経験'])

    # ---- 工数の集計
    r += 1
    ws.cell(row=r, column=1, value='工数の集計（分）').font = TITLE_F
    r += 1
    hdr = ['作業区分', '作業時間', '待ち時間', '合計', '件数']
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font, c.fill, c.border = HDR_F, HDR_FILL, BORDER
    first = r + 1
    rng = f'$E$4:$E${3 + BLANK_ROWS}'
    for kind in WORK_KINDS:
        r += 1
        ws.cell(row=r, column=1, value=kind)
        ws.cell(row=r, column=2, value=f'=SUMIF(工数!{rng},A{r},工数!$G$4:$G${3 + BLANK_ROWS})')
        ws.cell(row=r, column=3, value=f'=SUMIF(工数!{rng},A{r},工数!$H$4:$H${3 + BLANK_ROWS})')
        ws.cell(row=r, column=4, value=f'=B{r}+C{r}')
        ws.cell(row=r, column=5, value=f'=COUNTIF(工数!{rng},A{r})')
    r += 1
    ws.cell(row=r, column=1, value='合計').font = HDR_F
    for col in 'BCDE':
        ws[f'{col}{r}'] = f'=SUM({col}{first}:{col}{r - 1})'
    for row in ws.iter_rows(min_row=first, max_row=r, max_col=5):
        for c in row:
            c.border = BORDER
            if c.column > 1:
                c.fill = CALC_FILL
    r += 1
    ws.cell(row=r, column=1, value='手戻り率（手戻り ÷ 作業時間の合計）')
    ws.cell(row=r, column=2, value=f'=IF(B{r - 1}=0,"",B{first + WORK_KINDS.index("手戻り")}/B{r - 1})')
    ws.cell(row=r, column=2).number_format = '0%'
    for col in (1, 2):
        ws.cell(row=r, column=col).border = BORDER
    ws.cell(row=r, column=2).fill = CALC_FILL

    # ---- 結果・構築方法・差分の件数
    def counts(r, heading, sheet, col, values, last):
        r += 2
        ws.cell(row=r, column=1, value=heading).font = TITLE_F
        for v in values:
            r += 1
            ws.cell(row=r, column=1, value=v).border = BORDER
            c = ws.cell(row=r, column=2, value=f'=COUNTIF({sheet}!${col}$4:${col}${last},A{r})')
            c.border, c.fill = BORDER, CALC_FILL
        return r

    last_items = 3 + len(DESIGN_ITEMS) + 20
    r = counts(r, '設計項目の結果（件）', '設計項目別結果', 'F', RESULTS, last_items)
    r = counts(r, '設計項目の構築方法（件）', '設計項目別結果', 'G', METHODS, last_items)
    r = counts(r, '差分・修正の区分（件）', '差分・修正', 'D', DIFF_KINDS, 3 + BLANK_ROWS)

    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 40
    for col in 'CDE':
        ws.column_dimensions[col].width = 12


def sheet_effort(wb):
    ws = wb.create_sheet('工数')
    title(ws, '工数', '1 作業 1 行。作業区分は 6 つから選ぶ。手戻りは初回の作業と分けて記録する。経過は開始・終了から自動計算。')
    headers = ['No', '日付', '開始', '終了', '作業区分', '経過(分)', '作業時間(分)', '待ち時間(分)', '内容', '備考']
    first, last = table(ws, 3, headers, [5, 12, 8, 8, 16, 9, 11, 11, 44, 30],
                        input_cols={2, 3, 4, 5, 7, 8, 9, 10})
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)
        ws.cell(row=r, column=2).number_format = 'yyyy/mm/dd'
        ws.cell(row=r, column=3).number_format = 'hh:mm'
        ws.cell(row=r, column=4).number_format = 'hh:mm'
        c = ws.cell(row=r, column=6, value=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(MOD(D{r}-C{r},1)*1440,0))')
        c.fill = CALC_FILL
    dv_list(ws, WORK_KINDS, f'E{first}:E{last}')


def sheet_items(wb):
    ws = wb.create_sheet('設計項目別結果')
    title(ws, '設計項目別結果（構築できた項目 / 構築できなかった項目）',
          '節番号は設計書ごとに異なるため「設計書の節」に実際の番号を書く。設計書に無い項目は結果を「−」にする。')
    headers = ['No', '区分', '設計項目', '主な設定内容（参考）', '設計書の節', '結果', '構築方法', '担当ロール / 手順', '確認方法', '備考']
    fixed = [[i + 1, a, b, c] for i, (a, b, c) in enumerate(DESIGN_ITEMS)]
    first, last = table(ws, 3, headers, [5, 7, 26, 32, 10, 7, 12, 20, 11, 40],
                        rows=len(DESIGN_ITEMS) + 20, fixed=fixed, input_cols={5, 6, 7, 8, 9, 10})
    for r in range(first + len(DESIGN_ITEMS), last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).fill = INPUT_FILL
    dv_list(ws, RESULTS, f'F{first}:F{last}')
    dv_list(ws, METHODS, f'G{first}:G{last}')
    dv_list(ws, ['突合せ', '証跡', '目視', '未確認'], f'I{first}:I{last}')


def sheet_diff(wb):
    ws = wb.create_sheet('差分・修正')
    title(ws, '差分・修正', '既存資産から変えた箇所を 1 件 1 行。手動構築の場合は記入不要。')
    headers = ['No', '設計項目', '内容', '区分', '修正したファイル', 'コミット', '工数(分)', '標準に取り込むか', '備考']
    first, last = table(ws, 3, headers, [5, 22, 44, 13, 34, 10, 9, 13, 30], input_cols=set(range(2, 10)))
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)
    dv_list(ws, DIFF_KINDS, f'D{first}:D{last}')
    dv_list(ws, ['要', '不要', '要検討'], f'H{first}:H{last}')


def sheet_steps(wb):
    ws = wb.create_sheet('実施手順')
    title(ws, '実施手順', '実際にたどった順に書く。次の人がこのシートだけで再現できることを目安にする。')
    headers = ['順', '作業', 'コマンド・操作', '所要(分)', '備考']
    first, last = table(ws, 3, headers, [5, 30, 60, 9, 34], rows=40, input_cols={2, 3, 4, 5})
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)


def sheet_issues(wb):
    ws = wb.create_sheet('課題')
    title(ws, '課題', '詰まった点・設計書の不備・改善案を 1 件 1 行。')
    headers = ['No', '分類', '内容', '影響', '対応案', '状態']
    first, last = table(ws, 3, headers, [5, 11, 50, 30, 40, 10], rows=30, input_cols={2, 3, 4, 5, 6})
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - first + 1)
    dv_list(ws, ['設計書', 'コード', '環境', '手順', 'ツール', 'その他'], f'B{first}:B{last}')
    dv_list(ws, ['未対応', '対応中', '完了', '保留'], f'F{first}:F{last}')


def sheet_guide(wb):
    """記入要領の表（作業区分・結果・構築方法・差分の区分）を md から転記する。"""
    ws = wb.create_sheet('記入要領')
    title(ws, '記入要領（抜粋）', f'全文は {SRC} を参照。')
    text = io.open(SRC, encoding='utf-8').read()
    wanted = {'作業区分', '記号', '構築方法', '区分'}
    r = 4
    for hdr, rows in md_tables(text):
        if hdr[0] not in wanted:
            continue
        for i, h in enumerate(hdr, start=1):
            c = ws.cell(row=r, column=i, value=h)
            c.font, c.fill, c.border, c.alignment = HDR_F, HDR_FILL, BORDER, WRAP
        for row in rows:
            r += 1
            for i, v in enumerate(row, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.border, c.alignment = BORDER, WRAP
        r += 2
    for i, w in enumerate([18, 50, 50], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = openpyxl.Workbook()
    ref = {}
    sheet_basic(wb, ref)
    sheet_effort(wb)
    sheet_items(wb)
    sheet_diff(wb)
    sheet_steps(wb)
    sheet_issues(wb)
    sheet_guide(wb)
    wb.save(DST)
    print('wrote', DST)


if __name__ == '__main__':
    main()
